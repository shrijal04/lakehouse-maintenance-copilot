from pathlib import Path
from datetime import datetime

from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Image,
)

from docx import Document
from docx.shared import Inches


class ReportExportService:

    def __init__(self):

        # backend/reports
        self.report_folder = (
            Path(__file__).resolve().parents[2] / "reports"
        )

        self.report_folder.mkdir(exist_ok=True)

        # backend/reports/charts
        self.chart_folder = (
            self.report_folder / "charts"
        )

    # --------------------------------------------------
    # Add charts to PDF
    # --------------------------------------------------

    def _add_pdf_charts(
        self,
        story,
    ):

        chart_files = [
            "health_score.png",
            "snapshots.png",
            "data_files.png",
            "average_file_size.png",
            "storage.png",
        ]

        styles = getSampleStyleSheet()

        story.append(
            Paragraph(
                "<br/><b>Charts</b><br/><br/>",
                styles["Heading2"],
            )
        )

        for chart in chart_files:

            chart_path = self.chart_folder / chart

            if chart_path.exists():

                story.append(
                    Paragraph(
                        chart.replace("_", " ").replace(".png", "").title(),
                        styles["Heading3"],
                    )
                )

                story.append(
                    Image(
                        str(chart_path),
                        width=6 * inch,
                        height=3 * inch,
                    )
                )

                story.append(
                    Paragraph(
                        "<br/>",
                        styles["Normal"],
                    )
                )

    # --------------------------------------------------
    # Add charts to DOCX
    # --------------------------------------------------

    def _add_docx_charts(
        self,
        document,
    ):

        chart_files = [
            "health_score.png",
            "snapshots.png",
            "data_files.png",
            "average_file_size.png",
            "storage.png",
        ]

        document.add_heading(
            "Charts",
            level=2,
        )

        for chart in chart_files:

            chart_path = self.chart_folder / chart

            if chart_path.exists():

                document.add_heading(
                    chart.replace("_", " ").replace(".png", "").title(),
                    level=3,
                )

                document.add_picture(
                    str(chart_path),
                    width=Inches(6),
                )

    # --------------------------------------------------
    # Create PDF Report
    # --------------------------------------------------

    def export_pdf(
        self,
        title: str,
        report_text: str,
    ) -> str:

        timestamp = datetime.now().strftime(
            "%Y%m%d_%H%M%S"
        )

        filename = (
            f"daily_report_{timestamp}.pdf"
        )

        filepath = self.report_folder / filename

        document = SimpleDocTemplate(
            str(filepath)
        )

        styles = getSampleStyleSheet()

        story = []

        story.append(
            Paragraph(
                f"<b>{title}</b>",
                styles["Title"],
            )
        )

        story.append(
            Paragraph(
                f"Generated: {datetime.now()}",
                styles["Normal"],
            )
        )

        story.append(
            Paragraph(
                "<br/><br/>",
                styles["Normal"],
            )
        )

        for line in report_text.split("\n"):

            if line.strip():

                story.append(
                    Paragraph(
                        line,
                        styles["BodyText"],
                    )
                )

        # Add charts
        self._add_pdf_charts(story)

        document.build(story)

        return str(filepath)

    # --------------------------------------------------
    # Create DOCX Report
    # --------------------------------------------------

    def export_docx(
        self,
        title: str,
        report_text: str,
    ) -> str:

        timestamp = datetime.now().strftime(
            "%Y%m%d_%H%M%S"
        )

        filename = (
            f"daily_report_{timestamp}.docx"
        )

        filepath = self.report_folder / filename

        document = Document()

        document.add_heading(
            title,
            level=1,
        )

        document.add_paragraph(
            f"Generated: {datetime.now()}"
        )

        document.add_paragraph("")

        for line in report_text.split("\n"):

            document.add_paragraph(line)

        # Add charts
        self._add_docx_charts(document)

        document.save(filepath)

        return str(filepath)